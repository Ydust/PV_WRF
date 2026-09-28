"""Extract published manufacturing impacts without treating service units as STC watts."""
from pathlib import Path
import csv, json, hashlib
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'input/literature'
OUT = ROOT / 'input/parameters'

def main():
    w = load_workbook(SRC/'Willis_2026_source_data.xlsx', data_only=True, read_only=True)
    years = np.array([2023,2024,2026,2028,2031,2034])
    efficiency = np.array([.226,.228,.230,.234,.238,.240])
    first_degradation = np.array([.02,.015,.01,.01,.01,.01])
    # SI Tables 5/11 and Note 1: reverse the source's own functional-unit scaling.
    divisor = 1000 * efficiency * .8 * (1-first_degradation)
    topcon = np.array([w['Fig 1'].cell(5,c).value for c in range(6,10)])
    current2023 = w['Fig 3'].cell(5,2).value
    norm = current2023/topcon.sum()
    freight = topcon[-1]*norm*divisor[0]
    records=[]
    for region, col in [('china',2),('india',4),('usa',6),('europe',8)]:
        for grid, c in [('static_grid',col),('future_grid',col+1)]:
            for i,y in enumerate(years):
                raw=w['Fig 3'].cell(5+i,c).value
                records.append(dict(source='Willis et al. 2026',doi='10.1038/s41467-026-69165-x',
                    manufacturing_region=region,grid_case=grid,year=int(y),source_sheet='Fig 3',source_cell=f'{chr(64+c)}{5+i}',
                    reported_factor=raw,source_area_per_functional_unit_m2=1/divisor[i],
                    module_and_reference_freight_kg_co2e_m2=raw*divisor[i],
                    module_excluding_reference_freight_kg_co2e_m2=raw*divisor[i]-freight if region=='china' else '',
                    freight_kg_co2e_m2=freight if region=='china' else '',
                    transport_destination='Central Europe; reference route, not a global trade matrix'))
    with (OUT/'manufacturing_literature_factors.csv').open('w',newline='',encoding='utf8') as f:
        writer=csv.DictWriter(f,fieldnames=records[0]);writer.writeheader();writer.writerows(records)
    bos={
        'mounting_kg_co2e_m2':660/(3/.151),
        'electrical_kg_co2e_kwp':120/3,
        'inverter_kg_co2e_kwp_per_installation':420/3/2,
        'inverter_life_years':15,
        'inverter_count_interpretation':'Two generations over 30 years; tested separately as a modelling assumption.',
        'source':'IEA-PVPS T12-05:2015, Table 6.1, current single-Si rooftop system; 15.1% module efficiency.',
        'url':'https://iea-pvps.org/wp-content/uploads/2020/01/Future-PV-LCA-IEA-PVPS-Task-12-March-2015.pdf',
        'status':'Legacy rooftop BOS proxy; not a measured 2025 global inventory.',
        'replacement_support':'Lin et al. 2026, DOI 10.1038/s41467-026-76424-4, SI Table 2: 15-year inverter lifetime and one replacement.',
        'freight_kg_co2e_m2':float(freight),
        'freight_status':'Willis China-to-Central-Europe module route, applied as a common benchmark; sensitivity 0 to 2 times.',
        'end_of_life':'No PV modules retire within 2025-2050 (28-year life). No advance recycling credit.',
        'retired_inverter_treatment':'No recycling credit; treatment not independently quantified in the disaggregated reference.',
    }
    (OUT/'balance_of_system.json').write_text(json.dumps(bos,indent=2),encoding='utf8')
    # Continuation ratios, not invented absolute manufacturing inventories.
    tw=load_workbook(SRC/'Tang_2025_Table_S14.xlsx',data_only=True,read_only=True)['Fig 3(b)']
    continuation=[]
    for col,name in [(2,'ssp2_no_climate_target'),(3,'ssp2_2c'),(4,'ssp2_1p5c')]:
        pairs=[(int(tw.cell(row,1).value),float(tw.cell(row,col).value),row) for row in range(3,34)]
        anchor=next(v for y,v,_ in pairs if y==2034)
        for y,v,row in pairs:
            if y>=2025:continuation.append(dict(path=name,year=y,grid_factor_kg_co2e_kwh=v,ratio_to_2034=v/anchor,
                source='Tang et al. 2025; Table S14, Fig 3(b)',source_cell=f'{chr(64+col)}{row}',doi='10.1016/j.isci.2025.111963',
                role='Manufacturing-only scenario sensitivity; not a jointly simulated global transition'))
    cn=[]
    for y in [2021,2030,2040,2050]:
        with (ROOT/f'input/grids/global_grid_feedback_{y}_high.csv').open(encoding='utf8') as stream:
            row=next(v for v in csv.DictReader(stream) if v['country_tag']=='china')
        cn.append(float(row['country_co2']))
    anchor=np.interp(2034,[2021,2030,2040,2050],cn)
    for y in range(2025,2051):
        ef=float(np.interp(y,[2021,2030,2040,2050],cn))
        continuation.append(dict(path='inherited_china_grid',year=y,grid_factor_kg_co2e_kwh=ef,ratio_to_2034=ef/anchor,
            source='Inherited China deployment-grid trajectory',source_cell='',doi='',
            role='Shares the inherited China grid decline; absolute manufacturing and displaced-grid boundaries remain distinct'))
    with (OUT/'manufacturing_grid_continuations.csv').open('w',newline='',encoding='utf8') as stream:
        writer=csv.DictWriter(stream,fieldnames=continuation[0]);writer.writeheader();writer.writerows(continuation)
    # Keep acquisition provenance for every local source, including sources not used numerically.
    sources=[]
    for p in sorted(SRC.iterdir()):
        if p.suffix not in ('.pdf','.xlsx','.html'):continue
        sources.append(dict(file=p.name,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),bytes=p.stat().st_size))
    (SRC/'file_manifest.json').write_text(json.dumps(sources,indent=2),encoding='utf8')
    print('Manufacturing factors and explicitly labelled BOS proxy extracted.',flush=True)

if __name__=='__main__':main()
