"""Recover a source-model electricity response and update rooftop BOS benchmarks.

The recovered response is a model-derived coefficient, not metered factory energy.
All fitted observations originate from one published LCA, not independent factories.
"""
from pathlib import Path
import json,hashlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

R=Path(__file__).resolve().parents[1]

def main():
    s=load_workbook(R/'input/literature/Willis_2026_source_data.xlsx',data_only=True,read_only=True)
    w=load_workbook(R/'input/literature/Willis_2026_inventory.xlsx',data_only=True,read_only=True)
    factors=pd.read_csv(R/'input/parameters/manufacturing_literature_factors.csv')
    years=[2023,2024,2026,2028,2031,2034]
    countries=['china','india','usa','europe']
    grid=np.array([[w['Calculation 3'].cell(56+i,j).value for j in [2,3,4,5]] for i in range(6)])
    impact=np.array([[float(factors.query('year==@y and manufacturing_region==@c and grid_case=="future_grid"').module_and_reference_freight_kg_co2e_m2.iloc[0]) for c in countries] for y in years])
    # M(country,year) = E(year)*I(country,year) + R(year) + route_offset(country).
    # China offset=0 identifies the model. R initially includes China's freight.
    X=np.zeros((24,15))
    for i in range(6):
        for c in range(4):
            X[i*4+c,i]=grid[i,c];X[i*4+c,6+i]=1
            if c:X[i*4+c,11+c]=1
    q,_,rank,_=np.linalg.lstsq(X,impact.ravel(),rcond=None)
    assert rank==15 and np.all(q[:12]>0)
    pred=(X@q).reshape(6,4)
    freight=json.loads((R/'input/parameters/balance_of_system.json').read_text())['freight_kg_co2e_m2']
    # Anchor the residual to each exact China source result, avoiding fit drift.
    residual=impact[:,0]-freight-q[:6]*grid[:,0]
    assert np.all(residual>0)
    response=pd.DataFrame(dict(year=years,effective_electricity_response_kwh_m2=q[:6],
        source_china_grid_kg_co2e_kwh=grid[:,0],residual_kg_co2e_m2=residual,
        module_kg_co2e_m2=impact[:,0]-freight,
        responsive_fraction=q[:6]*grid[:,0]/(impact[:,0]-freight)))
    response.to_csv(R/'input/parameters/manufacturing_response.csv',index=False)
    observations=[];loo=[]
    for i,y in enumerate(years):
        for c,country in enumerate(countries):
            row=i*4+c;mask=np.arange(24)!=row
            b,_,rank_loo,_=np.linalg.lstsq(X[mask],impact.ravel()[mask],rcond=None)
            assert rank_loo==15
            loo.append(dict(year=y,country=country,held_out_error_kg_m2=float(X[row]@b-impact[i,c]),
                derived_2034_fraction=float(b[5]*grid[5,0]/(impact[5,0]-freight))))
            observations.append(dict(year=y,country=country,source_grid_cell=f'{chr(66+c)}{56+i}',
                source_impact_cell=f'{chr(67+2*c)}{5+i}',grid_kg_co2e_kwh=grid[i,c],
                source_kg_co2e_m2=impact[i,c],reconstructed_kg_co2e_m2=pred[i,c]))
    pd.DataFrame(observations).to_csv(R/'output/manufacturing_response_reconstruction.csv',index=False)
    pd.DataFrame(loo).to_csv(R/'output/manufacturing_response_holdouts.csv',index=False)
    # Static-China series was excluded from response fitting. Its six observations
    # share one unknown, constant grid intensity; use them as a consistency check.
    static=factors.query('manufacturing_region=="china" and grid_case=="static_grid"').sort_values('year').module_and_reference_freight_kg_co2e_m2.to_numpy()
    static_grid=float(np.dot(q[:6],static-q[6:12])/np.dot(q[:6],q[:6]))
    static_err=q[:6]*static_grid+q[6:12]-static
    checks=dict(full_rank=int(rank),condition_number=float(np.linalg.cond(X)),
        fit_rmse_kg_co2e_m2=float(np.sqrt(np.mean((pred-impact)**2))),
        maximum_held_out_error_kg_co2e_m2=max(abs(v['held_out_error_kg_m2']) for v in loo),
        maximum_static_series_error_kg_co2e_m2=float(abs(static_err).max()),
        implied_constant_static_grid_kg_co2e_kwh=static_grid,
        derived_2034_fraction=float(response.responsive_fraction.iloc[-1]),
        caveat='A response reconstructed from one source LCA, not measured factory electricity or territorial emissions. Residual includes all processes not responsive to the source grid change.',
        equations='M_c,y = E_y I_c,y + R_y + T_c; T_China=0. Exact China source observations anchor the final residual.')
    assert checks['maximum_held_out_error_kg_co2e_m2']<.01
    assert checks['maximum_static_series_error_kg_co2e_m2']<.001
    (R/'output/checks/manufacturing_response_validation.json').write_text(json.dumps(checks,indent=2))
    # Source Fig. 6 changes the China manufacturing sub-grid at fixed 2023
    # technology and freight. Recover relative grid burdens, not PV g/kWh.
    denom=float(1/factors.query('manufacturing_region=="china" and grid_case=="static_grid" and year==2023').source_area_per_functional_unit_m2.iloc[0])
    routes=[]
    for name,row in [('china_southwest',4),('china_average',5),('china_northeast',6)]:
        value=float(s['Fig 6'].cell(row,3).value)*denom
        implied=static_grid+(value-static[0])/q[0]
        assert implied>0
        routes.append(dict(route=name,source_cell=f'Fig 6!C{row}',
            source_2023_module_plus_freight_kg_co2e_m2=value,
            implied_grid_kg_co2e_kwh=implied,relative_grid_multiplier=implied/static_grid,
            status='Derived source-LCA sub-grid contrast; holding the ratio through 2050 is a scenario assumption, not an observed procurement forecast.'))
    pd.DataFrame(routes).to_csv(R/'input/parameters/manufacturing_subgrid_scenarios.csv',index=False)
    # Explicit printed values in IEA-PVPS 2024 slide 9 (2023 update).
    # 976 kWh/kWp/year already includes degradation: do not degrade it again.
    lifetime_yield=976*30
    bos=dict(source='IEA-PVPS 2024 slides, 2023 data update, slide 9 (printed component labels)',
        url='https://iea-pvps.org/wp-content/uploads/2024/05/Slides_IEA-PVPS-T12_Fact-Sheet-update-2023_v2.0.pdf',
        reference_system='3 kWp mono-Si slanted rooftop, Europe',
        reference_lifetime_yield_kwh_kwp=lifetime_yield,
        bos_manufacturing_g_co2e_kwh=5.1,inverter_manufacturing_g_co2e_kwh=9.9,
        combined_bos_kg_co2e_kwp=5.1/1000*lifetime_yield,
        inverter_kg_co2e_kwp_per_installation=9.9/1000*lifetime_yield/2,
        inverter_life_years=15,reference_module_life_years=30,
        replacement_basis='Two inverter generations follow the explicitly stated 30-year system and 15-year inverter lives.',
        component_boundary='Combined BOS is retained as one category; no unreported mounting/cable split is inferred.',
        scaling='Published combined BOS manufacturing burden is normalized to STC capacity. This is a reference-design scaling assumption, not a global equipment inventory.',
        rounding='Component labels have 0.1 g/kWh precision; half-unit rounding intervals are checked separately.',
        excluded_source_terms='PV module manufacture is replaced by Willis cohorts. Source residual 0.6 g/kWh includes other life-cycle stages and is not assigned wholesale to production. No future module disposal is booked inside 2025-2050.',
        scope='Published rooftop benchmark, not an observed global procurement mix. BOS manufacturing remains fixed unless separately varied.')
    (R/'input/parameters/balance_of_system_2024.json').write_text(json.dumps(bos,indent=2))
    manifest=[dict(file=p.name,bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest())
        for p in sorted((R/'input/literature').iterdir()) if p.suffix in ['.pdf','.xlsx','.html','.zip']]
    (R/'input/literature/file_manifest.json').write_text(json.dumps(manifest,indent=2))
    print('Source response recovered:',checks)
    print('Updated BOS and inverter kg/kWp:',bos['combined_bos_kg_co2e_kwp'],bos['inverter_kg_co2e_kwp_per_installation'])

if __name__=='__main__':main()
