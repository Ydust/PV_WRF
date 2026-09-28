"""Check cohort conservation, carbon balances and exported figure data."""
from pathlib import Path
from dataclasses import replace
import json
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from cohort_accounting import account,updated_cases,manufacturing_path

ROOT=Path(__file__).resolve().parents[1]

def main():
    raw=np.load(ROOT/'output/reference_75.npy');case=updated_cases()[3]
    result=account(raw,case);annual=result['service']
    np.testing.assert_allclose(annual[:,:,14],annual[:,:,8]-annual[:,:,9]-annual[:,:,10]-annual[:,:,11]+annual[:,:,12]-annual[:,:,13],atol=1e-12)
    for key in result['component_events']:
        np.testing.assert_allclose(result['component_events'][key].sum(1),result['component_service'][key].sum(1),rtol=2e-12,atol=1e-12)
    for i in list(range(9))+list(range(10,14))+list(range(15,20)):
        np.testing.assert_array_equal(annual[:,:,i],raw[:,:,i])
    single=np.zeros_like(raw);single[:,:,0]=raw[:,0,0,None];single[:,0,19]=raw[:,0,19]
    dynamic=account(single,case);frozen=account(single,replace(case,hold_manufacturing_grid_2025=True))
    np.testing.assert_array_equal(dynamic['component_events']['module'],frozen['component_events']['module'])
    np.testing.assert_array_equal(dynamic['component_service']['module'],frozen['component_service']['module'])
    np.testing.assert_allclose(annual[:,:,14].sum(),117.61190473366243,rtol=0,atol=1e-9)
    assert (annual[:,:,14].sum(1)<0).sum()==6
    np.testing.assert_allclose(-np.minimum(annual[:,:,14].sum(1),0).sum()*1000,75.64575417020752,rtol=0,atol=1e-8)
    low=manufacturing_path(updated_cases()[-2])[0];high=manufacturing_path(updated_cases()[-1])[0]
    ref=manufacturing_path(case)[0];assert np.all(low<=ref) and np.all(ref<=high)
    matched=pd.read_csv(ROOT/'output/Main_Figure_Data/matched_scenarios.csv')
    assert matched.sample_id.nunique()==2049
    assert abs(matched.delivered_twh/matched.target_twh-1).max()<1e-10
    book=load_workbook(ROOT/'output/Source_Data/Source_Data_Main_Figures.xlsx',read_only=True,data_only=True)
    paths=sorted((ROOT/'output/Source_Data/CSV').glob('*.csv'))
    assert len(paths)==17 and set(book.sheetnames)=={p.stem for p in paths}
    for p in paths:
        expected=pd.DataFrame(list(book[p.stem].values)[1:],columns=list(next(book[p.stem].values)))
        actual=pd.read_csv(p)
        pd.testing.assert_frame_equal(actual,expected,check_dtype=False,check_exact=False,rtol=1e-10,atol=1e-8)
    book.close()
    result={'passed':True,'net_mitigation_gt':float(annual[:,:,14].sum()),'net_emitting_countries':6,'source_data_sheets':17}
    (ROOT/'output/checks/verification.json').write_text(json.dumps(result,indent=2))
    print(json.dumps(result,indent=2))

if __name__=='__main__':main()
